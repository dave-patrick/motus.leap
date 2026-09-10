import csv
import os
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

input_csv = r"C:\Users\davem\Downloads\maintenance-export-2026-08-09T23-46-23.csv"
output_xlsx = r"C:\Users\davem\Downloads\maintenance_queue_segmented.xlsx"

print(f"Reading CSV from {input_csv}...")
with open(input_csv, mode="r", encoding="utf-8-sig") as f:
    reader = csv.DictReader(f)
    rows = list(reader)

print(f"Total rows read: {len(rows)}")

# Segment data
exact_dups = [r for r in rows if r.get("queue_type") == "duplicate" and r.get("exact_duplicate", "").lower() == "true"]
cross_dups = [r for r in rows if r.get("queue_type") == "duplicate" and r.get("exact_duplicate", "").lower() != "true"]
misplaced = [r for r in rows if r.get("queue_type") == "misplaced"]

print(f"Exact Duplicates: {len(exact_dups)}")
print(f"Cross-Listed Duplicates: {len(cross_dups)}")
print(f"Misplaced Items: {len(misplaced)}")

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Styles
font_title = Font(name="Calibri", size=16, bold=True, color="1F4E79")
font_subtitle = Font(name="Calibri", size=11, italic=True, color="595959")
font_sec_hdr = Font(name="Calibri", size=13, bold=True, color="1F4E79")
font_tbl_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=11, bold=True)
font_regular = Font(name="Calibri", size=11)
font_link = Font(name="Calibri", size=11, color="0563C1", underline="single")

fill_exact_hdr = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Dark Blue
fill_cross_hdr = PatternFill(start_color="006666", end_color="006666", fill_type="solid") # Dark Teal
fill_misplaced_hdr = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid") # Steel Blue
fill_summary_hdr = PatternFill(start_color="333333", end_color="333333", fill_type="solid") # Dark Gray
fill_zebra = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

# ---------------------------------------------------------
# Tab 1: Summary
# ---------------------------------------------------------
ws_sum = wb.create_sheet(title="Summary")
ws_sum.views.sheetView[0].showGridLines = True

ws_sum["A1"] = "Maintenance Queue Overview & Summary"
ws_sum["A1"].font = font_title
ws_sum["A2"] = "Segmented report generated from maintenance queue export"
ws_sum["A2"].font = font_subtitle

# Table 1: Actions Breakdown
ws_sum["A4"] = "Queue Actions Breakdown"
ws_sum["A4"].font = font_sec_hdr

headers_t1 = ["Queue Segment", "Row Count", "% of Total", "Description"]
for col_idx, h in enumerate(headers_t1, 1):
    cell = ws_sum.cell(row=5, column=col_idx, value=h)
    cell.font = font_tbl_hdr
    cell.fill = fill_summary_hdr
    cell.alignment = Alignment(horizontal="center", vertical="center")

total_count = len(rows)
t1_data = [
    ("Exact Duplicates", len(exact_dups), f"{len(exact_dups)/total_count:.1%}", "Identical videos appearing multiple times within playlists"),
    ("Cross-Listed Duplicates", len(cross_dups), f"{len(cross_dups)/total_count:.1%}", "Videos present in multiple different playlists"),
    ("Misplaced Items", len(misplaced), f"{len(misplaced)/total_count:.1%}", "Videos categorized in non-optimal target playlists"),
    ("Total Queue Items", total_count, "100.0%", "Complete maintenance queue size")
]

for r_idx, row_data in enumerate(t1_data, 6):
    is_total = (r_idx == 9)
    for c_idx, val in enumerate(row_data, 1):
        cell = ws_sum.cell(row=r_idx, column=c_idx, value=val)
        cell.font = font_bold if is_total else font_regular
        cell.border = thin_border
        if c_idx in (2, 3):
            cell.alignment = Alignment(horizontal="right")

# Table 2: Top Target Playlists for Misplaced Items
ws_sum["A12"] = "Top Target Destinations (Misplaced Items)"
ws_sum["A12"].font = font_sec_hdr

headers_t2 = ["Target Playlist (`to_playlist`)", "Misplaced Count", "% of Misplaced"]
for col_idx, h in enumerate(headers_t2, 1):
    cell = ws_sum.cell(row=13, column=col_idx, value=h)
    cell.font = font_tbl_hdr
    cell.fill = fill_summary_hdr
    cell.alignment = Alignment(horizontal="center", vertical="center")

to_pl_counts = Counter(r['to_playlist'] for r in misplaced if r.get('to_playlist'))
top_to_pl = to_pl_counts.most_common(15)

for r_idx, (pl_name, count) in enumerate(top_to_pl, 14):
    cell_a = ws_sum.cell(row=r_idx, column=1, value=pl_name)
    cell_b = ws_sum.cell(row=r_idx, column=2, value=count)
    cell_c = ws_sum.cell(row=r_idx, column=3, value=f"{count/len(misplaced):.1%}")
    cell_a.font = font_regular
    cell_b.font = font_regular
    cell_c.font = font_regular
    cell_a.border = thin_border
    cell_b.border = thin_border
    cell_c.border = thin_border
    cell_b.alignment = Alignment(horizontal="right")
    cell_c.alignment = Alignment(horizontal="right")

# ---------------------------------------------------------
# Helper function for data tabs
# ---------------------------------------------------------
def create_data_tab(ws_name, data_rows, headers, fill_hdr, is_dup_tab=False):
    ws = wb.create_sheet(title=ws_name)
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    
    # Header row
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = font_tbl_hdr
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 24

    for r_idx, row_item in enumerate(data_rows, 2):
        vid = row_item.get("video_id", "")
        yt_url = f"https://www.youtube.com/watch?v={vid}" if vid else ""
        
        if is_dup_tab:
            copy_cnt = row_item.get("copy_count", "")
            try:
                copy_cnt = int(copy_cnt)
            except ValueError:
                pass
            row_values = [
                vid,
                yt_url,
                row_item.get("title", ""),
                row_item.get("channel", ""),
                row_item.get("from_playlist", ""),
                copy_cnt,
                row_item.get("all_playlists", "")
            ]
        else:
            row_values = [
                vid,
                yt_url,
                row_item.get("title", ""),
                row_item.get("channel", ""),
                row_item.get("from_playlist", ""),
                row_item.get("to_playlist", ""),
                row_item.get("all_playlists", "")
            ]
            
        for c_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            cell.font = font_regular
            
            if c_idx == 2 and val:
                cell.value = val
                cell.font = font_link
            elif c_idx == 6 and is_dup_tab:
                cell.value = val
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.value = val
                
            if r_idx % 2 == 1:
                cell.fill = fill_zebra

    # Add auto filter
    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col_letter}{len(data_rows)+1}"

    # Auto adjust column widths
    for col in ws.columns:
        col_idx = col[0].column
        max_len = max(len(str(cell.value or '')) for cell in col[:100]) # sample first 100 rows
        header_len = len(str(ws.cell(row=1, column=col_idx).value or ''))
        width = max(max_len + 3, header_len + 3)
        # Cap widths for very long titles/URLs
        if col_idx in (2, 3, 7): # URL, title, all_playlists
            width = min(width, 50)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(width, 12)

# Create the data tabs
dup_headers = ["Video ID", "YouTube URL", "Video Title", "Channel", "From Playlist", "Copy Count", "All Playlists"]
misplaced_headers = ["Video ID", "YouTube URL", "Video Title", "Channel", "From Playlist", "To Playlist", "All Playlists"]

print("Creating 'Exact Duplicates' tab...")
create_data_tab("Exact Duplicates", exact_dups, dup_headers, fill_exact_hdr, is_dup_tab=True)

print("Creating 'Cross-Listed Duplicates' tab...")
create_data_tab("Cross-Listed Duplicates", cross_dups, dup_headers, fill_cross_hdr, is_dup_tab=True)

print("Creating 'Misplaced Items' tab...")
create_data_tab("Misplaced Items", misplaced, misplaced_headers, fill_misplaced_hdr, is_dup_tab=False)

# Auto adjust summary tab column widths
for col in ws_sum.columns:
    col_idx = col[0].column
    max_len = max(len(str(cell.value or '')) for cell in col)
    ws_sum.column_dimensions[get_column_letter(col_idx)].width = max(max_len + 4, 15)

print(f"Saving Excel workbook to {output_xlsx}...")
wb.save(output_xlsx)
print("Workbook saved successfully!")
